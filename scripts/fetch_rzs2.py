#!/usr/bin/env python3
"""Phase 2 RZS indicators -> data/clean/
  starost.csv     average age (latest year)
  prirastaj.csv   live births, deaths and natural increase per 1,000 inhabitants (latest year)
  vrtici.csv      children not admitted to public kindergartens (capacity), children enrolled, private share
  turizam.csv     tourist overnight stays, last 12 months
  kriminal.csv    adults convicted for violent + property crimes by place of offence, per 10,000 inhabitants
                  (from the yearbook "Opštine i regioni u Republici Srbiji", Excel)
"""
import csv
import io
import json
import re
import sys
import urllib.request
from collections import defaultdict
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from fetch_rzs import CITY_IDS, MONTHS, download, lower_keys  # noqa: E402
from geo_units import units  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
CLEAN = ROOT / "data" / "clean"
RAW = ROOT / "data" / "raw"
UA = "Mozilla/5.0 (compatible; gdeziveti-bot/1.0; +https://gdeziveti.rs)"

_C2L = {"А": "A", "Б": "B", "В": "V", "Г": "G", "Д": "D", "Ђ": "Đ", "Е": "E", "Ж": "Ž", "З": "Z", "И": "I", "Ј": "J",
        "К": "K", "Л": "L", "Љ": "Lj", "М": "M", "Н": "N", "Њ": "Nj", "О": "O", "П": "P", "Р": "R", "С": "S", "Т": "T",
        "Ћ": "Ć", "У": "U", "Ф": "F", "Х": "H", "Ц": "C", "Ч": "Č", "Џ": "Dž", "Ш": "Š"}
_C2L.update({k.lower(): v.lower() for k, v in _C2L.items()})


def c2l(s):
    return "".join(_C2L.get(ch, ch) for ch in s)


def meta_update(updates):
    p = CLEAN / "meta.json"
    meta = json.loads(p.read_text()) if p.exists() else {}
    meta.update(updates)
    p.write_text(json.dumps(meta, ensure_ascii=False, indent=1), encoding="utf-8")


def id_map_from(rows):
    names = {}
    for r in rows:
        names.setdefault(r["nter"].strip(), r["idter"])
    out = {}
    for u in units():
        n = u["name"]
        out[n] = CITY_IDS.get(n) or names.get(n)
    return out


def latest(rows, filt):
    rs = [r for r in rows if filt(r) and r["vrednost"] not in ("", None)]
    y = max(r["god"] for r in rs)
    return y, {r["idter"]: float(r["vrednost"]) for r in rs if r["god"] == y}


def write(name, header, rows):
    with (CLEAN / name).open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


def main():
    pop = {}
    with (CLEAN / "stanovnistvo.csv").open(encoding="utf-8") as f:
        for r in csv.DictReader(f):
            pop[r["naziv"]] = float(r["pop_procena"] or 0) or None

    # ---- average age ----
    a = lower_keys(download("180711IND01"))
    ids = id_map_from(a)
    y_age, age = latest(a, lambda r: r["idpol"] == "0" and r["idtipnaselja"] == "0")
    write("starost.csv", ["naziv", "prosecna_starost"], [[n, age[i]] for n, i in ids.items() if i in age])

    # ---- natural increase ----
    b = lower_keys(download("180304IND07"))
    d = lower_keys(download("180304IND08"))
    y_b, br = latest(b, lambda r: True)
    y_d, dr = latest(d, lambda r: r["god"] == y_b)
    write("prirastaj.csv", ["naziv", "rodjeni_na_1000", "umrli_na_1000", "prirastaj_na_1000"],
          [[n, br[i], dr[i], round(br[i] - dr[i], 1)] for n, i in ids.items() if i in br and i in dr])

    # ---- kindergartens ----
    k1 = lower_keys(download("110103010IND01"))
    k2 = lower_keys(download("110103010IND02"))
    y_k, enrolled = latest(k1, lambda r: r["idsvojina"] == "0")
    _, private = latest(k1, lambda r: r["idsvojina"] == "1" and r["god"] == y_k)
    _, rejected = latest(k2, lambda r: r["idsvojina"] == "0" and r["god"] == y_k)
    rows = []
    for n, i in ids.items():
        if i in enrolled and enrolled[i] > 0:
            rej = rejected.get(i, 0.0)
            rows.append([n, int(rej), int(enrolled[i]), round(rej / (enrolled[i] + rej) * 100, 1),
                         round(private.get(i, 0.0) / enrolled[i] * 100, 1)])
    write("vrtici.csv", ["naziv", "nisu_primljeni", "upisani", "odbijeni_pct", "privatni_pct"], rows)

    # ---- tourism (last 12 months) ----
    t = lower_keys(download("220205IND02"))
    t = [r for r in t if r["idturisti"] == "0" and r["vrednost"] not in ("", None)]
    periods = sorted({(r["god"], r["mes"]) for r in t})[-12:]
    nights = defaultdict(float)
    for r in t:
        if (r["god"], r["mes"]) in periods:
            nights[r["idter"]] += float(r["vrednost"])
    write("turizam.csv", ["naziv", "nocenja_12m", "nocenja_po_stanovniku"],
          [[n, int(nights[i]), round(nights[i] / pop[n], 1) if pop.get(n) else ""] for n, i in ids.items() if i in nights])
    (g0, m0), (g1, m1) = periods[0], periods[-1]

    # ---- crime (yearbook) ----
    crime_period = fetch_crime(pop)

    meta_update({
        "starost": {"period": f"procena {y_age}.", "year": int(y_age), "verified": True},
        "prirastaj": {"period": f"{y_b}. godina", "year": int(y_b), "verified": True},
        "vrtici": {"period": f"{y_k}. godina", "year": int(y_k), "verified": True},
        "turizam": {"period": f"{MONTHS[int(m0)-1]} {g0} – {MONTHS[int(m1)-1]} {g1}.", "year": int(g1), "verified": True},
        **({"kriminal": crime_period} if crime_period else {}),
    })
    print("age", y_age, "| births", y_b, "| kindergartens", y_k, "| tourism", periods[0], periods[-1], "| crime", crime_period)


def fetch_crime(pop):
    import openpyxl
    wb_bytes, y_pub = None, None
    for y in (date.today().year, date.today().year - 1, date.today().year - 2):
        for n in range(13040, 13080):
            url = f"https://publikacije.stat.gov.rs/G{y}/Xls/G{y}{n}.xlsx"
            try:
                data = urllib.request.urlopen(urllib.request.Request(url, headers={"User-Agent": UA}), timeout=60).read()
            except Exception:  # noqa: BLE001
                continue
            if data[:2] == b"PK":
                wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
                sheet = next((s for s in wb.sheetnames if s.startswith("19.")), None)
                first = next(wb[sheet].iter_rows(max_row=1, values_only=True), ()) if sheet else ()
                if sheet and "ОСУЂЕН" in " ".join(str(c) for c in first if c):
                    wb_bytes, y_pub = data, y
                    break
        if wb_bytes:
            break
    if not wb_bytes:
        print("crime: yearbook not found; keeping previous file")
        return None
    (RAW / "opstine_i_regioni.xlsx").write_bytes(wb_bytes)
    wb = openpyxl.load_workbook(io.BytesIO(wb_bytes), read_only=True)
    ws = wb[next(s for s in wb.sheetnames if s.startswith("19."))]
    rows = list(ws.iter_rows(values_only=True))
    title = str(rows[0][1] or "")
    m = re.search(r"(20\d\d)", title)
    data_year = int(m.group(1)) if m else y_pub - 1
    names = {u["name"] for u in units()}
    lower_names = {x.lower(): x for x in names}
    found = {}
    for r in rows[5:]:
        code = str(r[0] or "").strip()
        raw = str(r[1] or "").strip()
        if code not in ("1", "I") or not raw:
            continue
        n = " ".join(c2l(raw).replace("Grad ", "").split())
        if n == "Palilula":
            n = "Palilula (Beograd)"
        n = lower_names.get(n.lower(), n)
        if n in names and n not in found:
            try:
                found[n] = (int(r[2] or 0), int(r[3] or 0), int(r[4] or 0))
            except (TypeError, ValueError):
                pass
    out = []
    for n, (total, life, prop) in found.items():
        if pop.get(n):
            out.append([n, total, life, prop, round((life + prop) / pop[n] * 10000, 1)])
    write("kriminal.csv", ["naziv", "osudjeni_ukupno", "protiv_zivota_i_tela", "protiv_imovine", "nasilni_imovinski_na_10000"], out)
    print("crime units:", len(out), "missing:", sorted(names - set(found)))
    return {"period": f"{data_year}. godina", "year": data_year, "verified": True}


if __name__ == "__main__":
    main()
